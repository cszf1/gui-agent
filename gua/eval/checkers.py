"""基于最终状态的客观判分（WebArena / OSWorld 的 execution-based evaluation 思路）。

判分器只在评测端运行，结果不会泄露给 agent。
- 文件类（所有桌面平台）：file_exists / file_contains / excel_cell / docx_contains / dir_count
- Web：web_text / web_js / web_url（需要 env，直接读 DOM，不依赖截图）
- Android：android_text（读 uiautomator 树）/ android_foreground
- Mock：mock_state
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Callable

CHECKERS: dict[str, Callable[..., tuple[bool, str]]] = {}
NEEDS_ENV: set[str] = set()


def checker(name: str, needs_env: bool = False):
    def deco(fn):
        CHECKERS[name] = fn
        if needs_env:
            NEEDS_ENV.add(name)
        return fn
    return deco


def _p(path: str) -> Path:
    import os
    return Path(os.path.expandvars(os.path.expanduser(path)))


@checker("file_exists")
def file_exists(path: str) -> tuple[bool, str]:
    p = _p(path)
    return p.exists(), f"{p} exists={p.exists()}"


@checker("file_contains")
def file_contains(path: str, text: str = "", regex: str = "", encoding: str = "utf-8") -> tuple[bool, str]:
    p = _p(path)
    if not p.exists():
        return False, f"{p} missing"
    s = p.read_text(encoding=encoding, errors="ignore")
    ok = (text in s) if text else bool(re.search(regex, s))
    return ok, f"{p} contains={ok}"


@checker("file_not_contains")
def file_not_contains(path: str, text: str, encoding: str = "utf-8") -> tuple[bool, str]:
    p = _p(path)
    if not p.exists():
        return False, f"{p} missing"
    ok = text not in p.read_text(encoding=encoding, errors="ignore")
    return ok, f"{p} not_contains={ok}"


@checker("excel_cell")
def excel_cell(path: str, cell: str, equals: Any = None, sheet: str | None = None,
               tol: float = 1e-6) -> tuple[bool, str]:
    import openpyxl
    p = _p(path)
    if not p.exists():
        return False, f"{p} missing"
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    v = ws[cell].value
    if isinstance(equals, (int, float)) and isinstance(v, (int, float)):
        ok = abs(v - equals) <= tol
    else:
        ok = str(v) == str(equals)
    return ok, f"{cell}={v!r} expected {equals!r}"


@checker("excel_has_chart")
def excel_has_chart(path: str, sheet: str | None = None) -> tuple[bool, str]:
    import openpyxl
    p = _p(path)
    if not p.exists():
        return False, f"{p} missing"
    wb = openpyxl.load_workbook(p)
    ws = wb[sheet] if sheet else wb.active
    n = len(getattr(ws, "_charts", []))
    return n > 0, f"charts={n}"


@checker("docx_contains")
def docx_contains(path: str, text: str) -> tuple[bool, str]:
    import docx
    p = _p(path)
    if not p.exists():
        return False, f"{p} missing"
    full = "\n".join(par.text for par in docx.Document(p).paragraphs)
    return text in full, f"docx contains={text in full}"


@checker("dir_count")
def dir_count(path: str, pattern: str = "*", equals: int = 0) -> tuple[bool, str]:
    n = len(list(_p(path).glob(pattern)))
    return n == equals, f"{n} files match {pattern}"


@checker("web_text", needs_env=True)
def web_text(env, text: str, selector: str = "body", absent: bool = False) -> tuple[bool, str]:
    got = env.eval_js(f"(() => {{ const n = document.querySelector({json.dumps(selector)}); return n ? n.innerText : null; }})()")
    present = got is not None and text in got
    ok = (not present) if absent else present
    return ok, f"{selector} {'lacks' if absent else 'contains'} {text!r}: {ok} (got {str(got)[:80]!r})"


@checker("web_js", needs_env=True)
def web_js(env, expr: str, equals: Any = True) -> tuple[bool, str]:
    got = env.eval_js(expr)
    return got == equals, f"js {expr[:60]!r} = {got!r}, expected {equals!r}"


@checker("web_url", needs_env=True)
def web_url(env, contains: str) -> tuple[bool, str]:
    url = env.page.url
    return contains in url, f"url={url!r} contains {contains!r}"


@checker("android_text", needs_env=True)
def android_text(env, text: str) -> tuple[bool, str]:
    obs = env.observe(with_elements=True)
    ok = text in obs.all_text()
    return ok, f"android screen contains {text!r}: {ok}"


@checker("android_foreground", needs_env=True)
def android_foreground(env, package: str) -> tuple[bool, str]:
    comp, pkg = env.foreground()
    return pkg == package, f"foreground={comp!r}"


@checker("mock_state", needs_env=True)
def mock_state(env, key: str, equals: Any = True) -> tuple[bool, str]:
    v = env.state.get(key)
    return v == equals, f"state[{key!r}]={v!r} expected {equals!r}"


def run_checks(specs: list[dict], env=None) -> tuple[bool, list[str]]:
    notes, all_ok = [], True
    for s in specs:
        s = dict(s)
        name = s.pop("type")
        fn = CHECKERS[name]
        try:
            ok, note = fn(env, **s) if name in NEEDS_ENV else fn(**s)
        except Exception as e:  # noqa: BLE001
            ok, note = False, f"{name} crashed: {type(e).__name__}: {e}"
        notes.append(("PASS " if ok else "FAIL ") + note)
        all_ok &= ok
    return all_ok, notes
